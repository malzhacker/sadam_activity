from pathlib import Path
from itertools import count

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Student_Learning_Hub_20_Word_Shapes_Portrait.docx"

# Namespaces used by modern, editable Word Processing Shapes (DrawingML).
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
EMU_PER_PT = 12700
SHAPE_IDS = count(1)

# Exactly 20 requested interfaces.
SCREENS = [
    ("Login", "Akaun", ["E-mel", "Kata laluan"], "LOG MASUK", "Sistem sahkan akaun", "Login berjaya!", "Dashboard dibuka", "auth"),
    ("Pendaftaran Pelajar", "Akaun", ["Nama / ID", "E-mel", "Kata laluan"], "DAFTAR", "Sistem semak dan simpan", "Pendaftaran berjaya", "Halaman login dipaparkan", "form"),
    ("Lupa Kata Laluan", "Akaun", ["E-mel berdaftar"], "HANTAR PAUTAN", "Sistem hantar e-mel reset", "Pautan telah dihantar", "Arahan semak e-mel", "auth"),
    ("Dashboard Pelajar", "Utama", ["Pilih kad aktiviti"], "BUKA", "Sistem ambil ringkasan", "Tugasan hampir tamat!", "Kursus dan kemajuan", "dashboard"),
    ("Profil Pelajar", "Akaun", ["Gambar", "Telefon", "Bio"], "SIMPAN", "Sistem kemas kini rekod", "Profil berjaya dikemas kini", "Profil baharu dipaparkan", "profile"),
    ("Senarai Kursus", "Kursus", ["Carian", "Kategori"], "CARI", "Sistem tapis kursus", "Kursus disimpan", "Senarai kursus sepadan", "cards"),
    ("Butiran Kursus", "Kursus", ["Pilih tab / modul"], "BUKA MODUL", "Sistem ambil kandungan", "Modul berjaya dibuka", "Info pensyarah dan modul", "detail"),
    ("Pendaftaran Kursus", "Kursus", ["Semester", "Seksyen", "Setuju syarat"], "DAFTAR KURSUS", "Semak kelayakan dan kapasiti", "Sahkan pendaftaran?", "Status: BERJAYA", "form"),
    ("Modul Pembelajaran", "Pembelajaran", ["Pilih topik"], "TANDA SELESAI", "Sistem rekod kemajuan", "Kemajuan direkodkan", "Modul seterusnya dibuka", "learning"),
    ("Bahan Pembelajaran", "Pembelajaran", ["Carian", "Jenis bahan"], "MUAT TURUN", "Sistem sediakan fail", "Fail berjaya dimuat turun", "Nota, slaid dan video", "table"),
    ("Senarai Tugasan", "Tugasan", ["Kursus", "Status"], "TAPIS", "Susun ikut tarikh akhir", "Tarikh akhir hari ini!", "Senarai dan status tugasan", "list"),
    ("Butiran Tugasan", "Tugasan", ["Pilih lampiran / rubrik"], "BUKA", "Sistem ambil arahan", "Lampiran dimuat turun", "Arahan dan markah penuh", "detail"),
    ("Penghantaran Tugasan", "Tugasan", ["Fail DOCX/PDF", "Komen"], "HANTAR", "Sistem semak dan upload", "Pasti mahu hantar?", "Resit penghantaran", "upload"),
    ("Keputusan Tugasan", "Tugasan", ["Pilih tugasan dinilai"], "LIHAT", "Sistem ambil keputusan", "Markah telah diterbitkan", "86/100 dan komen", "result"),
    ("Senarai Kuiz", "Kuiz", ["Kursus", "Kuiz tersedia"], "MULAKAN", "Sistem semak masa", "Masa 15 minit bermula", "Soalan kuiz dibuka", "list"),
    ("Menjawab Kuiz", "Kuiz", ["Pilih jawapan A-D"], "HANTAR KUIZ", "Sistem simpan jawapan", "Masih ada soalan kosong!", "Jawapan direkodkan", "quiz"),
    ("Keputusan Kuiz", "Kuiz", ["Pilih semak jawapan"], "SEMAK", "Sistem kira skor", "Kuiz selesai: 80%", "8 betul dan 2 salah", "result"),
    ("Forum Perbincangan", "Komunikasi", ["Tajuk", "Kandungan topik"], "TERBIT", "Sistem simpan topik", "Topik berjaya diterbitkan", "Topik dan balasan", "forum"),
    ("Pengumuman", "Komunikasi", ["Pilih pengumuman"], "BUKA", "Sistem tanda dibaca", "Ditanda sebagai dibaca", "Butiran pengumuman", "announcement"),
    ("Notifikasi", "Komunikasi", ["Kategori", "Status baca"], "TANDA DIBACA", "Sistem kemas kini status", "Semua notifikasi dibaca", "Notifikasi terkini", "notifications"),
]


def tag(namespace, local):
    return f"{{{namespace}}}{local}"


def pt(value):
    return str(int(value * EMU_PER_PT))


def add_text_paragraph(container, value, size=9, bold=False, color="222222", align="center"):
    p = etree.SubElement(container, tag(W, "p"))
    ppr = etree.SubElement(p, tag(W, "pPr"))
    jc = etree.SubElement(ppr, tag(W, "jc"))
    jc.set(tag(W, "val"), align)
    spacing = etree.SubElement(ppr, tag(W, "spacing"))
    spacing.set(tag(W, "before"), "0")
    spacing.set(tag(W, "after"), "0")
    spacing.set(tag(W, "line"), "210")
    spacing.set(tag(W, "lineRule"), "auto")
    run = etree.SubElement(p, tag(W, "r"))
    rpr = etree.SubElement(run, tag(W, "rPr"))
    fonts = etree.SubElement(rpr, tag(W, "rFonts"))
    fonts.set(tag(W, "ascii"), "Calibri")
    fonts.set(tag(W, "hAnsi"), "Calibri")
    fonts.set(tag(W, "eastAsia"), "Calibri")
    fonts.set(tag(W, "cs"), "Calibri")
    if bold:
        etree.SubElement(rpr, tag(W, "b"))
    color_el = etree.SubElement(rpr, tag(W, "color"))
    color_el.set(tag(W, "val"), color)
    sz = etree.SubElement(rpr, tag(W, "sz"))
    sz.set(tag(W, "val"), str(size * 2))
    sz_cs = etree.SubElement(rpr, tag(W, "szCs"))
    sz_cs.set(tag(W, "val"), str(size * 2))
    text_el = etree.SubElement(run, tag(W, "t"))
    text_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text_el.text = value


def add_shape(paragraph, x, y, width, height, text="", geometry="rect", fill="FFFFFF",
              line="222222", line_width=1.1, font_size=9, bold=False, text_color="222222",
              align="center", dashed=False, flip_h=False, flip_v=False, arrow=False,
              name="Word Shape"):
    """Add an independently editable native Word DrawingML shape anchored to the page."""
    shape_id = next(SHAPE_IDS)
    run = paragraph.add_run()
    drawing = etree.SubElement(run._r, tag(W, "drawing"))
    anchor = etree.SubElement(drawing, tag(WP, "anchor"), {
        "distT": "0", "distB": "0", "distL": "0", "distR": "0",
        "simplePos": "0", "relativeHeight": str(251658240 + shape_id),
        "behindDoc": "0", "locked": "0", "layoutInCell": "1", "allowOverlap": "1"
    })
    etree.SubElement(anchor, tag(WP, "simplePos"), {"x": "0", "y": "0"})
    pos_h = etree.SubElement(anchor, tag(WP, "positionH"), {"relativeFrom": "page"})
    etree.SubElement(pos_h, tag(WP, "posOffset")).text = pt(x)
    pos_v = etree.SubElement(anchor, tag(WP, "positionV"), {"relativeFrom": "page"})
    etree.SubElement(pos_v, tag(WP, "posOffset")).text = pt(y)
    etree.SubElement(anchor, tag(WP, "extent"), {"cx": pt(width), "cy": pt(height)})
    etree.SubElement(anchor, tag(WP, "effectExtent"), {"l": "0", "t": "0", "r": "0", "b": "0"})
    etree.SubElement(anchor, tag(WP, "wrapNone"))
    etree.SubElement(anchor, tag(WP, "docPr"), {"id": str(shape_id), "name": f"{name} {shape_id}"})
    nv = etree.SubElement(anchor, tag(WP, "cNvGraphicFramePr"))
    etree.SubElement(nv, tag(A, "graphicFrameLocks"), {"noChangeAspect": "0"})

    graphic = etree.SubElement(anchor, tag(A, "graphic"))
    data = etree.SubElement(graphic, tag(A, "graphicData"), {
        "uri": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
    })
    wsp = etree.SubElement(data, tag(WPS, "wsp"))
    cnv = etree.SubElement(wsp, tag(WPS, "cNvSpPr"))
    if text:
        cnv.set("txBox", "1")
    sppr = etree.SubElement(wsp, tag(WPS, "spPr"))
    xfrm_attrs = {}
    if flip_h:
        xfrm_attrs["flipH"] = "1"
    if flip_v:
        xfrm_attrs["flipV"] = "1"
    xfrm = etree.SubElement(sppr, tag(A, "xfrm"), xfrm_attrs)
    etree.SubElement(xfrm, tag(A, "off"), {"x": "0", "y": "0"})
    etree.SubElement(xfrm, tag(A, "ext"), {"cx": pt(width), "cy": pt(height)})
    geom = etree.SubElement(sppr, tag(A, "prstGeom"), {"prst": geometry})
    etree.SubElement(geom, tag(A, "avLst"))

    if fill is None:
        etree.SubElement(sppr, tag(A, "noFill"))
    else:
        solid = etree.SubElement(sppr, tag(A, "solidFill"))
        etree.SubElement(solid, tag(A, "srgbClr"), {"val": fill})

    ln = etree.SubElement(sppr, tag(A, "ln"), {"w": str(int(line_width * EMU_PER_PT))})
    if line is None:
        etree.SubElement(ln, tag(A, "noFill"))
    else:
        line_fill = etree.SubElement(ln, tag(A, "solidFill"))
        etree.SubElement(line_fill, tag(A, "srgbClr"), {"val": line})
        if dashed:
            etree.SubElement(ln, tag(A, "prstDash"), {"val": "dash"})
        if arrow:
            etree.SubElement(ln, tag(A, "tailEnd"), {"type": "triangle", "w": "med", "len": "med"})

    if text:
        txbx = etree.SubElement(wsp, tag(WPS, "txbx"))
        content = etree.SubElement(txbx, tag(W, "txbxContent"))
        lines = text.split("\n")
        for value in lines:
            add_text_paragraph(content, value, font_size, bold, text_color, align)
        body = etree.SubElement(wsp, tag(WPS, "bodyPr"), {
            "rot": "0", "spcFirstLastPara": "0", "vertOverflow": "overflow",
            "horzOverflow": "overflow", "vert": "horz", "wrap": "square",
            "lIns": pt(3), "tIns": pt(2), "rIns": pt(3), "bIns": pt(2),
            "numCol": "1", "spcCol": "0", "rtlCol": "0", "fromWordArt": "0",
            "anchor": "ctr", "anchorCtr": "0", "forceAA": "0", "compatLnSpc": "1"
        })
        etree.SubElement(body, tag(A, "spAutoFit"))
    else:
        body = etree.SubElement(wsp, tag(WPS, "bodyPr"), {"rot": "0"})
        etree.SubElement(body, tag(A, "spAutoFit"))
    return shape_id


def textbox(p, x, y, w, h, value, size=9, bold=False, align="left", color="222222"):
    return add_shape(p, x, y, w, h, value, "rect", None, None, font_size=size,
                     bold=bold, text_color=color, align=align, name="Text Box")


def rectangle(p, x, y, w, h, text="", rounded=False, fill="FFFFFF", line="222222",
              size=9, bold=False, dashed=False, align="center"):
    return add_shape(p, x, y, w, h, text, "roundRect" if rounded else "rect", fill, line,
                     font_size=size, bold=bold, dashed=dashed, align=align,
                     name="Rounded Rectangle" if rounded else "Rectangle")


def oval(p, x, y, w, h, text="", fill="FFFFFF", line="222222", size=9, bold=False):
    return add_shape(p, x, y, w, h, text, "ellipse", fill, line, font_size=size, bold=bold,
                     name="Oval")


def line_shape(p, x, y, w, h=1, arrow=False, flip_h=False, dashed=False):
    return add_shape(p, x, y, max(w, 1), max(h, 1), "", "line", None, "333333", 1.3,
                     dashed=dashed, flip_h=flip_h, arrow=arrow, name="Arrow" if arrow else "Line")


def page_title(doc, number, title, module):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"Interface {number:02d}: {title}")
    r.font.name = "Calibri"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(35, 35, 35)
    r2 = p.add_run(f"    Modul: {module}")
    r2.font.name = "Calibri"
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor(95, 95, 95)
    anchor_p = doc.add_paragraph()
    anchor_p.paragraph_format.space_after = Pt(0)
    anchor_p.paragraph_format.line_spacing = Pt(1)
    return anchor_p


def panel(p, x, y, number, label):
    rectangle(p, x, y, 270, 330, "", False, "FFFFFF", "555555", dashed=False)
    oval(p, x + 8, y + 8, 22, 22, str(number), "FFFFFF", "555555", 9, False)
    textbox(p, x + 36, y + 8, 224, 22, label, 9, False, "left")


def browser(p, x, y, w, h, title="Student Learning Hub"):
    rectangle(p, x, y, w, h, "", True, "FFFFFF", "777777")
    line_shape(p, x, y + 20, w, 1)
    oval(p, x + 7, y + 6, 7, 7, "", "FFFFFF", "888888", 5)
    oval(p, x + 19, y + 6, 7, 7, "", "FFFFFF", "888888", 5)
    oval(p, x + 31, y + 6, 7, 7, "", "FFFFFF", "888888", 5)
    textbox(p, x + 47, y + 3, w - 54, 14, title, 7, False, "left", "666666")


def input_panel(p, x, y, fields, action, layout):
    sx, sy, sw, sh = x + 12, y + 39, 246, 274
    browser(p, sx, sy, sw, sh)
    content_x = sx + (52 if layout not in ("auth", "form") else 12)
    if layout not in ("auth", "form"):
        rectangle(p, sx + 5, sy + 27, 40, sh - 34, "SLH\n\nHome\nCourse\nTask", False,
                  "F2F2F2", "999999", 7, False)
    textbox(p, content_x, sy + 28, sw - (content_x - sx) - 9, 18, "Input pengguna", 8, False, "left")
    shown = fields[:3]
    usable = sw - (content_x - sx) - 17
    label_width = 72 if layout not in ("auth", "form") else 82
    for idx, field in enumerate(shown):
        fy = sy + 55 + idx * 43
        textbox(p, content_x, fy, label_width, 18, field, 7, False, "left")
        rectangle(p, content_x, fy + 18, usable, 22, "________________", True,
                  "FFFFFF", "888888", 7, False, align="left")
    button_y = sy + 62 + len(shown) * 43
    rectangle(p, content_x, min(button_y, sy + 205), min(usable, 130), 25,
              action, True, "E7E7E7", "555555", 8, False)
    textbox(p, content_x, sy + 235, usable, 18, "Klik butang selepas isi maklumat", 7, False, "left", "666666")


def process_panel(p, x, y, process_text):
    sx, sy, sw, sh = x + 12, y + 39, 246, 274
    browser(p, sx, sy, sw, sh)
    textbox(p, sx + 12, sy + 29, sw - 24, 18, "Proses dalam sistem", 8, False, "center")
    center_x = sx + sw / 2
    stages = [(sy + 57, "Terima input"), (sy + 119, "Semak maklumat"), (sy + 181, "Simpan / paparkan")]
    for stage_y, label in stages:
        rectangle(p, center_x - 67, stage_y, 134, 34, label, True, "FFFFFF", "666666", 8, False)
    line_shape(p, center_x, sy + 92, 1, 26, True)
    line_shape(p, center_x, sy + 154, 1, 26, True)
    textbox(p, sx + 14, sy + 228, sw - 28, 28, process_text, 8, False, "center")


def popup_panel(p, x, y, message):
    sx, sy, sw, sh = x + 12, y + 39, 246, 274
    browser(p, sx, sy, sw, sh)
    for row in range(5):
        rectangle(p, sx + 11, sy + 31 + row * 33, sw - 22, 21, "", True,
                  "F8F8F8", "BBBBBB", dashed=True)
    modal_x, modal_y, modal_w, modal_h = sx + 27, sy + 69, 192, 137
    rectangle(p, modal_x, modal_y, modal_w, modal_h, "", True, "FFFFFF", "444444")
    textbox(p, modal_x + 12, modal_y + 11, modal_w - 24, 18, "Popup message", 9, False, "left")
    textbox(p, modal_x + 12, modal_y + 36, modal_w - 24, 45, message, 8, False, "left")
    confirm = any(word in message.lower() for word in ("sahkan", "pasti", "bermula", "kosong"))
    if confirm:
        rectangle(p, modal_x + 12, modal_y + 96, 65, 24, "Batal", True, "FFFFFF", "666666", 8, False)
        rectangle(p, modal_x + modal_w - 77, modal_y + 96, 65, 24, "Ya", True, "E7E7E7", "555555", 8, False)
    else:
        rectangle(p, modal_x + modal_w - 77, modal_y + 96, 65, 24, "OK", True, "E7E7E7", "555555", 8, False)
    textbox(p, sx + 12, sy + 238, sw - 24, 19, "Mesej keluar selepas proses", 7, False, "left", "666666")


def output_panel(p, x, y, output, layout):
    sx, sy, sw, sh = x + 12, y + 39, 246, 274
    browser(p, sx, sy, sw, sh)
    rectangle(p, sx + 5, sy + 27, 40, sh - 34, "SLH\n\nHome\nCourse\nTask", False,
              "F2F2F2", "999999", 7, False)
    cx = sx + 53
    cw = sw - 62
    textbox(p, cx, sy + 29, cw, 18, "Output sistem", 8, False, "left")
    if layout in ("dashboard", "cards"):
        for j in range(3):
            rectangle(p, cx + j * 61, sy + 56, 55, 52, f"Kad {j + 1}\n____", True,
                      "FFFFFF", "777777", 7, False)
        rectangle(p, cx, sy + 120, cw - 1, 75, "Graf / kemajuan\n\n /  /  /", False,
                  "FFFFFF", "888888", 7, False)
    elif layout == "result":
        oval(p, cx + 3, sy + 59, 73, 73, "86\n/100" if "86" in output else "80%",
             "FFFFFF", "555555", 12, False)
        for j in range(3):
            rectangle(p, cx + 86, sy + 59 + j * 36, cw - 92, 24, "______", True,
                      "FFFFFF", "888888", 7, False)
    elif layout == "profile":
        oval(p, cx + 5, sy + 59, 61, 61, "Foto", "FFFFFF", "777777", 8, False)
        for j in range(3):
            rectangle(p, cx + 78, sy + 57 + j * 36, cw - 84, 24, "______", True,
                      "FFFFFF", "888888", 7, False)
    elif layout == "quiz":
        for j, option in enumerate(["A", "B", "C", "D"]):
            oval(p, cx, sy + 54 + j * 34, 19, 19, option, "FFFFFF", "777777", 7, False)
            line_shape(p, cx + 28, sy + 63 + j * 34, cw - 34, 1)
    else:
        for j in range(5):
            rectangle(p, cx, sy + 52 + j * 33, cw - 1, 24, f"{j + 1}.  ______________", True,
                      "FFFFFF", "888888", 7, False, align="left")
    textbox(p, cx, sy + 232, cw, 24, output, 8, False, "left")


def flow_arrows(p):
    line_shape(p, 290, 222, 15, 1, True)
    textbox(p, 280, 202, 35, 14, "Next", 7, False, "center", "666666")
    line_shape(p, 442, 390, 1, 15, True)
    textbox(p, 448, 390, 42, 14, "Popup", 7, False, "left", "666666")
    line_shape(p, 290, 572, 15, 1, True, True)
    textbox(p, 280, 578, 35, 14, "Hasil", 7, False, "center", "666666")


def draw_storyboard(p, screen):
    title, module, fields, action, process_text, popup, output, layout = screen
    coords = [(18, 58), (307, 58), (18, 405), (307, 405)]
    labels = [
        "Input - pengguna isi atau pilih",
        "Proses - tindakan dalam sistem",
        "Popup - mesej kepada pengguna",
        "Output - hasil sistem",
    ]
    for number, ((x, y), label) in enumerate(zip(coords, labels), 1):
        panel(p, x, y, number, label)
    input_panel(p, *coords[0], fields, action, layout)
    process_panel(p, *coords[1], process_text)
    popup_panel(p, *coords[2], popup)
    output_panel(p, *coords[3], output, layout)
    flow_arrows(p)
    textbox(p, 22, 748, 360, 18, "Aliran: Input  ->  Proses  ->  Popup  ->  Output", 8, False, "left", "666666")
    textbox(p, 385, 748, 188, 18, "Boleh diedit dalam Word", 8, False, "right", "666666")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def setup_section(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.4)
    section.right_margin = Inches(0.4)
    section.header_distance = Inches(0.12)
    section.footer_distance = Inches(0.18)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Muka surat ")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def heading(doc, title, subtitle=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.font.name = "Calibri"
    r.font.bold = True
    r.font.size = Pt(22)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(subtitle)
        r2.font.name = "Calibri"
        r2.font.size = Pt(10)
        r2.font.color.rgb = RGBColor(90, 90, 90)


def build_document():
    doc = Document()
    setup_section(doc.sections[0])
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    # Plain cover intentionally resembles a normal student-made Word assignment.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(120)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Student Learning Hub System")
    r.font.name = "Calibri"
    r.font.bold = True
    r.font.size = Pt(22)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("20 Lakaran Interface Menggunakan Word Shapes")
    r2.font.name = "Calibri"
    r2.font.bold = False
    r2.font.size = Pt(16)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.add_run("Input  ->  Proses  ->  Popup Message  ->  Output")
    doc.add_paragraph()
    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row, label in zip(table.rows, ["Nama Ahli 1", "Nama Ahli 2", "Nama Ahli 3", "Kursus / Kelas"]):
        row.cells[0].text = label
        row.cells[1].text = "____________________________________________"
        row.cells[0].paragraphs[0].runs[0].bold = True
        set_cell_shading(row.cells[0], "E7E7E7")
        for cell in row.cells:
            set_cell_margins(cell, 110, 140, 110, 140)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(18)
    nr = note.add_run("Nota: Semua lakaran dibina menggunakan rectangle, arrow, line dan text box Microsoft Word. Tiada gambar digunakan.")
    nr.font.name = "Calibri"
    nr.font.size = Pt(9)
    doc.add_page_break()

    # Interface index only; the task-allocation section is intentionally omitted.
    heading(doc, "Senarai 20 Interface")
    index = doc.add_table(rows=1, cols=3)
    index.style = "Table Grid"
    index.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(["Bil.", "Nama interface", "Modul"]):
        index.rows[0].cells[i].text = value
        index.rows[0].cells[i].paragraphs[0].runs[0].bold = True
        set_cell_shading(index.rows[0].cells[i], "E7E7E7")
    for number, screen in enumerate(SCREENS, 1):
        cells = index.add_row().cells
        for i, value in enumerate((f"{number:02d}", screen[0], screen[1])):
            cells[i].text = value
            set_cell_margins(cells[i], 35, 70, 35, 70)
    doc.add_page_break()

    # 20 pages of editable native Word shapes.
    for number, screen in enumerate(SCREENS, 1):
        anchor_p = page_title(doc, number, screen[0], screen[1])
        draw_storyboard(anchor_p, screen)
        if number < len(SCREENS):
            doc.add_page_break()

    for section in doc.sections:
        setup_section(section)
        add_page_number(section.footer.paragraphs[0])

    props = doc.core_properties
    props.title = "Student Learning Hub - 20 Portrait Word Shapes Storyboards"
    props.subject = "A4 portrait, editable Microsoft Word shapes: input, process, popup and output"
    props.author = "Student Learning Hub Group"
    props.keywords = "Word shapes, portrait, Calibri, storyboard, interface, editable, no images"
    doc.save(OUTPUT)
    print(f"CREATED={OUTPUT}")
    print(f"INTERFACES={len(SCREENS)}")
    print(f"SIZE_BYTES={OUTPUT.stat().st_size}")


if __name__ == "__main__":
    build_document()
