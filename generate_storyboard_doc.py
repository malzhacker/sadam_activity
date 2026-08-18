from pathlib import Path
import random
import re
import textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
PREVIEW_DIR = ROOT / "storyboard_previews"
PREVIEW_DIR.mkdir(exist_ok=True)
OUTPUT = ROOT / "Student_Learning_Hub_20_Storyboard_Sketch.docx"

W, H = 1500, 840
PAPER = "#fffdf7"
INK = "#282828"
PENCIL = "#575757"
FAINT = "#c9c5ba"
LIGHT = "#ece8dd"
WHITE = "#ffffff"

FONT = "/usr/share/fonts/google-noto/NotoSans-Regular.ttf"
FONT_BOLD = "/usr/share/fonts/google-noto/NotoSans-Bold.ttf"


def f(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT, size)


F12, F14, F16, F18, F22, F28, F34 = [f(s) for s in (12, 14, 16, 18, 22, 28, 34)]
B12, B14, B16, B18, B22, B28, B34 = [f(s, True) for s in (12, 14, 16, 18, 22, 28, 34)]

# Exactly 20 interfaces. Each item describes one four-frame storyboard.
SCREENS = [
    ("Login", "Akaun", ["E-mel", "Kata laluan"], "Tekan LOG MASUK", "Login berjaya!", "Dashboard dibuka", "auth"),
    ("Pendaftaran Pelajar", "Akaun", ["Nama / ID", "E-mel", "Kata laluan"], "Tekan DAFTAR", "Akaun berjaya didaftarkan", "Halaman login dipaparkan", "form"),
    ("Lupa Kata Laluan", "Akaun", ["E-mel berdaftar"], "Hantar pautan reset", "Pautan telah dihantar", "Arahan semak e-mel", "auth"),
    ("Dashboard Pelajar", "Utama", ["Pilih kad aktiviti"], "Sistem ambil ringkasan", "Tugasan hampir tamat!", "Kursus & kemajuan dipaparkan", "dashboard"),
    ("Profil Pelajar", "Akaun", ["Gambar", "Telefon", "Bio"], "Simpan perubahan", "Profil berjaya dikemas kini", "Profil baharu dipaparkan", "profile"),
    ("Senarai Kursus", "Kursus", ["Carian", "Penapis kategori"], "Cari / tapis kursus", "Kursus disimpan", "Senarai kursus sepadan", "cards"),
    ("Butiran Kursus", "Kursus", ["Pilih tab / modul"], "Buka kandungan kursus", "Modul berjaya dibuka", "Info pensyarah & modul", "detail"),
    ("Pendaftaran Kursus", "Kursus", ["Semester", "Seksyen", "Setuju syarat"], "Daftar kursus", "Sahkan pendaftaran?", "Status: BERJAYA", "form"),
    ("Modul Pembelajaran", "Pembelajaran", ["Pilih topik"], "Tanda topik selesai", "Kemajuan direkodkan", "Modul seterusnya dibuka", "learning"),
    ("Bahan Pembelajaran", "Pembelajaran", ["Carian", "Jenis bahan"], "Pilih MUAT TURUN", "Fail berjaya dimuat turun", "Nota / slaid / video", "table"),
    ("Senarai Tugasan", "Tugasan", ["Kursus", "Status tugasan"], "Tapis ikut tarikh", "Tarikh akhir hari ini!", "Senarai & status tugasan", "list"),
    ("Butiran Tugasan", "Tugasan", ["Pilih lampiran / rubrik"], "Buka arahan tugasan", "Lampiran dimuat turun", "Arahan & markah penuh", "detail"),
    ("Penghantaran Tugasan", "Tugasan", ["Fail DOCX/PDF", "Komen"], "Tekan HANTAR", "Pasti mahu hantar?", "Resit penghantaran", "upload"),
    ("Keputusan Tugasan", "Tugasan", ["Pilih tugasan dinilai"], "Sistem ambil keputusan", "Markah telah diterbitkan", "86/100 + komen", "result"),
    ("Senarai Kuiz", "Kuiz", ["Kursus", "Kuiz tersedia"], "Tekan MULAKAN", "Masa 15 minit bermula", "Soalan kuiz dibuka", "list"),
    ("Menjawab Kuiz", "Kuiz", ["Pilih jawapan A–D"], "Simpan & hantar jawapan", "Masih ada soalan kosong!", "Jawapan direkodkan", "quiz"),
    ("Keputusan Kuiz", "Kuiz", ["Pilih semak jawapan"], "Sistem kira skor", "Kuiz selesai: 80%", "8 betul / 2 salah", "result"),
    ("Forum Perbincangan", "Komunikasi", ["Tajuk", "Kandungan topik"], "Terbit topik", "Topik berjaya diterbitkan", "Topik & balasan dipaparkan", "forum"),
    ("Pengumuman", "Komunikasi", ["Pilih pengumuman"], "Buka & tanda dibaca", "Ditanda sebagai dibaca", "Butiran pengumuman", "announcement"),
    ("Notifikasi", "Komunikasi", ["Kategori", "Status baca"], "Tanda semua dibaca", "Semua notifikasi dibaca", "Senarai notifikasi terkini", "notifications"),
]


def text(draw, xy, value, font=F14, fill=INK, anchor=None):
    draw.text(xy, value, font=font, fill=fill, anchor=anchor)


def wrapped(draw, xy, value, width, font=F14, fill=PENCIL, spacing=4):
    chars = max(8, int(width / (font.size * 0.58)))
    draw.multiline_text(xy, "\n".join(textwrap.wrap(value, chars)), font=font, fill=fill, spacing=spacing)


def rough_line(draw, points, fill=INK, width=2, passes=2, jitter=1.5):
    for _ in range(passes):
        moved = [(int(x + random.uniform(-jitter, jitter)), int(y + random.uniform(-jitter, jitter))) for x, y in points]
        draw.line(moved, fill=fill, width=width)


def rough_box(draw, box, fill=None, outline=INK, width=2, radius=0):
    x1, y1, x2, y2 = box
    if fill:
        if radius:
            draw.rounded_rectangle(box, radius=radius, fill=fill)
        else:
            draw.rectangle(box, fill=fill)
    if radius:
        # Two slightly offset rounded outlines create a hand-drawn look.
        for _ in range(2):
            j = random.randint(-1, 1)
            draw.rounded_rectangle((x1 + j, y1, x2, y2 + j), radius=radius, outline=outline, width=width)
    else:
        rough_line(draw, [(x1, y1), (x2, y1)], outline, width)
        rough_line(draw, [(x2, y1), (x2, y2)], outline, width)
        rough_line(draw, [(x2, y2), (x1, y2)], outline, width)
        rough_line(draw, [(x1, y2), (x1, y1)], outline, width)


def scribble_fill(draw, box, density=10):
    x1, y1, x2, y2 = box
    for y in range(y1 + 3, y2, density):
        rough_line(draw, [(x1 + 3, y), (x2 - 3, y + random.randint(-2, 2))], FAINT, 1, 1, 1)


def arrow(draw, start, end, label=""):
    x1, y1 = start
    x2, y2 = end
    rough_line(draw, [(x1, y1), (x2, y2)], PENCIL, 3, 2, 2)
    angle = __import__("math").atan2(y2 - y1, x2 - x1)
    for offset in (2.55, -2.55):
        a = angle + offset
        p = (x2 + 13 * __import__("math").cos(a), y2 + 13 * __import__("math").sin(a))
        rough_line(draw, [(x2, y2), p], PENCIL, 3, 2, 1)
    if label:
        text(draw, ((x1 + x2) / 2, (y1 + y2) / 2 - 14), label, B12, PENCIL, "mm")


def numbered_circle(draw, x, y, number):
    draw.ellipse((x - 18, y - 18, x + 18, y + 18), fill=PAPER, outline=INK, width=2)
    text(draw, (x, y), str(number), B16, INK, "mm")


def browser_frame(draw, box, title="Student Learning Hub"):
    x1, y1, x2, y2 = box
    rough_box(draw, box, WHITE, INK, 2, 7)
    rough_line(draw, [(x1, y1 + 28), (x2, y1 + 28)], INK, 2)
    for i in range(3):
        draw.ellipse((x1 + 12 + i * 17, y1 + 10, x1 + 21 + i * 17, y1 + 19), outline=PENCIL, width=1)
    text(draw, (x1 + 75, y1 + 14), title, F12, PENCIL, "lm")
    return x1 + 12, y1 + 40, x2 - 12, y2 - 10


def input_field(draw, box, label, value="________________"):
    x1, y1, x2, y2 = box
    text(draw, (x1, y1), label, B12, INK)
    rough_box(draw, (x1, y1 + 18, x2, y2), WHITE, PENCIL, 1, 4)
    text(draw, (x1 + 9, y1 + 33), value, F12, PENCIL, "lm")


def button(draw, box, label, filled=False):
    rough_box(draw, box, INK if filled else WHITE, INK, 2, 4)
    x1, y1, x2, y2 = box
    text(draw, ((x1 + x2) / 2, (y1 + y2) / 2), label, B12, WHITE if filled else INK, "mm")


def tiny_sidebar(draw, box, active=0):
    x1, y1, x2, y2 = box
    draw.rectangle(box, fill=LIGHT)
    text(draw, (x1 + 9, y1 + 14), "SLH", B12, INK)
    for i in range(5):
        yy = y1 + 39 + i * 24
        if i == active:
            scribble_fill(draw, (x1 + 7, yy - 5, x2 - 7, yy + 13), 6)
        rough_line(draw, [(x1 + 12, yy), (x2 - 15, yy)], PENCIL, 2, 1, 1)


def draw_form_scene(draw, box, fields, action, click=False, layout="form"):
    x1, y1, x2, y2 = browser_frame(draw, box)
    if layout not in ("auth", "form"):
        tiny_sidebar(draw, (x1, y1, x1 + 82, y2), 1)
        x1 += 100
    text(draw, (x1 + 12, y1 + 6), "INPUT", B14, INK)
    max_fields = min(3, len(fields))
    for idx, label in enumerate(fields[:max_fields]):
        yy = y1 + 35 + idx * 44
        input_field(draw, (x1 + 12, yy, x2 - 18, yy + 36), label)
    by = min(y2 - 38, y1 + 43 + max_fields * 44)
    button(draw, (x2 - 250, by, x2 - 18, by + 30), action.upper()[:28], True)
    if click:
        text(draw, (x2 - 125, by + 43), "^ klik sini", B12, PENCIL, "mm")
        arrow(draw, (x2 - 120, by + 37), (x2 - 135, by + 23))


def draw_process_scene(draw, box, action, layout):
    x1, y1, x2, y2 = browser_frame(draw, box)
    tiny_sidebar(draw, (x1, y1, x1 + 82, y2), 2 if layout in ("upload", "result", "list") else 1)
    text(draw, (x1 + 108, y1 + 8), "PROSES SISTEM", B14, INK)
    cy = (y1 + y2) // 2
    stages = [(x1 + 145, "INPUT"), (x1 + 285, "SEMAK"), (x1 + 425, "SIMPAN")]
    for sx, label in stages:
        draw.ellipse((sx - 35, cy - 35, sx + 35, cy + 35), fill=WHITE, outline=INK, width=2)
        scribble_fill(draw, (sx - 25, cy - 5, sx + 25, cy + 7), 5)
        text(draw, (sx, cy + 52), label, B12, INK, "mm")
    arrow(draw, (stages[0][0] + 38, cy), (stages[1][0] - 38, cy))
    arrow(draw, (stages[1][0] + 38, cy), (stages[2][0] - 38, cy))
    wrapped(draw, (x1 + 105, y2 - 36), action, x2 - x1 - 120, F12, PENCIL)


def draw_popup_scene(draw, box, message, confirm=False):
    x1, y1, x2, y2 = browser_frame(draw, box)
    tiny_sidebar(draw, (x1, y1, x1 + 78, y2), 1)
    # Faint background content.
    for j in range(4):
        rough_box(draw, (x1 + 105, y1 + 18 + j * 35, x2 - 22, y1 + 42 + j * 35), None, FAINT, 1, 3)
    mw = min(380, x2 - x1 - 150)
    mx1 = x1 + 115
    my1 = y1 + 34
    rough_box(draw, (mx1, my1, mx1 + mw, my1 + 135), PAPER, INK, 3, 7)
    text(draw, (mx1 + 18, my1 + 20), "POPUP MESSAGE", B14, INK)
    wrapped(draw, (mx1 + 18, my1 + 48), message, mw - 36, F14, INK)
    if confirm:
        button(draw, (mx1 + 18, my1 + 96, mx1 + 122, my1 + 124), "BATAL")
        button(draw, (mx1 + mw - 128, my1 + 96, mx1 + mw - 18, my1 + 124), "YA", True)
    else:
        button(draw, (mx1 + mw - 118, my1 + 96, mx1 + mw - 18, my1 + 124), "OK", True)


def draw_output_scene(draw, box, output, layout):
    x1, y1, x2, y2 = browser_frame(draw, box)
    tiny_sidebar(draw, (x1, y1, x1 + 82, y2), 0 if layout == "dashboard" else 1)
    ax1 = x1 + 100
    text(draw, (ax1, y1 + 6), "OUTPUT", B14, INK)
    if layout in ("dashboard", "cards"):
        for j in range(3):
            cx = ax1 + j * 145
            rough_box(draw, (cx, y1 + 37, cx + 125, y1 + 108), WHITE, PENCIL, 1, 5)
            scribble_fill(draw, (cx + 12, y1 + 52, cx + 56, y1 + 68), 5)
            rough_line(draw, [(cx + 12, y1 + 88), (cx + 105, y1 + 88)], PENCIL, 2, 1)
        rough_box(draw, (ax1, y1 + 125, x2 - 18, y2 - 8), WHITE, PENCIL, 1, 4)
        rough_line(draw, [(ax1 + 15, y2 - 30), (ax1 + 85, y2 - 55), (ax1 + 165, y2 - 42), (ax1 + 245, y2 - 75), (x2 - 35, y2 - 95)], INK, 3)
    elif layout in ("result",):
        draw.ellipse((ax1 + 20, y1 + 42, ax1 + 145, y1 + 167), outline=INK, width=4)
        text(draw, (ax1 + 82, y1 + 103), re.findall(r"\d+", output)[0] if re.findall(r"\d+", output) else "OK", B28, INK, "mm")
        for j in range(4):
            yy = y1 + 40 + j * 34
            rough_line(draw, [(ax1 + 190, yy), (x2 - 25, yy)], PENCIL, 2, 1)
    elif layout == "quiz":
        for j in range(4):
            yy = y1 + 36 + j * 36
            draw.ellipse((ax1, yy, ax1 + 17, yy + 17), outline=INK, width=2)
            rough_line(draw, [(ax1 + 30, yy + 8), (x2 - 30, yy + 8)], PENCIL, 2, 1)
    elif layout == "notifications":
        for j in range(3):
            yy = y1 + 34 + j * 39
            draw.ellipse((ax1, yy, ax1 + 18, yy + 18), outline=INK, width=1)
            rough_line(draw, [(ax1 + 30, yy + 5), (x2 - 25, yy + 5)], PENCIL, 2, 1)
            rough_line(draw, [(ax1 + 30, yy + 16), (x2 - 95, yy + 16)], FAINT, 1, 1)
    else:
        for j in range(4):
            yy = y1 + 34 + j * 31
            rough_box(draw, (ax1, yy, x2 - 22, yy + 23), WHITE, PENCIL, 1, 2)
            rough_line(draw, [(ax1 + 12, yy + 11), (x2 - 80, yy + 11)], PENCIL, 2, 1)
    wrapped(draw, (ax1, y2 - 30), output, x2 - ax1 - 20, B12, INK)


def draw_storyboard(index, item):
    random.seed(index * 7919)
    title, module, fields, action, popup, output, layout = item
    image = Image.new("RGB", (W, H), PAPER)
    draw = ImageDraw.Draw(image)

    # Faint paper guide lines make it look like a photographed storyboard sheet.
    for y in range(18, H, 28):
        draw.line((0, y, W, y), fill="#f3efe5", width=1)
    draw.line((34, 0, 34, H), fill="#ead8d4", width=2)

    text(draw, (55, 24), f"STORYBOARD #{index:02d}", B18, PENCIL)
    text(draw, (55, 54), title.upper(), B28, INK)
    text(draw, (1440, 34), "STUDENT LEARNING HUB", B14, PENCIL, "rm")
    text(draw, (1440, 60), f"Modul: {module}  |  Lakaran awal", F12, PENCIL, "rm")

    panels = [
        (45, 105, 735, 420),
        (765, 105, 1455, 420),
        (45, 460, 735, 785),
        (765, 460, 1455, 785),
    ]
    labels = ["INPUT — pengguna isi / pilih", "PROSES — sistem jalankan tindakan", "POPUP — mesej kepada pengguna", "OUTPUT — hasil dipaparkan"]

    for n, (panel, label) in enumerate(zip(panels, labels), 1):
        rough_box(draw, panel, PAPER, INK, 2)
        numbered_circle(draw, panel[0] + 28, panel[1] + 28, n)
        text(draw, (panel[0] + 57, panel[1] + 27), label, B14, INK, "lm")

    inner = [(p[0] + 25, p[1] + 55, p[2] - 25, p[3] - 25) for p in panels]
    draw_form_scene(draw, inner[0], fields, action, False, layout)
    draw_process_scene(draw, inner[1], action, layout)
    draw_popup_scene(draw, inner[2], popup, any(k in popup.lower() for k in ("sahkan", "pasti", "bermula", "kosong")))
    draw_output_scene(draw, inner[3], output, layout)

    arrow(draw, (735, 260), (765, 260), "kemudian")
    arrow(draw, (1110, 420), (1110, 460), "mesej")
    arrow(draw, (765, 625), (735, 625), "akhirnya")

    path = PREVIEW_DIR / f"{index:02d}_{title.lower().replace(' ', '_')}_storyboard.png"
    image.save(path, quality=95)
    return path


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_margins(cell, top=60, start=80, bottom=60, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def setup_section(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = Inches(0.3)
    section.bottom_margin = Inches(0.3)
    section.left_margin = Inches(0.42)
    section.right_margin = Inches(0.42)
    section.header_distance = Inches(0.15)
    section.footer_distance = Inches(0.15)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Muka surat ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(90, 90, 90)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def centered_heading(doc, title, subtitle=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.font.name = "Arial"
    r.font.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(40, 40, 40)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(8)
        r2 = p2.add_run(subtitle)
        r2.font.name = "Arial"
        r2.font.size = Pt(10)
        r2.font.color.rgb = RGBColor(85, 85, 85)


def build_document(paths):
    doc = Document()
    setup_section(doc.sections[0])
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor(40, 40, 40)

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(55)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("STUDENT LEARNING HUB SYSTEM")
    r.font.name = "Arial"
    r.font.bold = True
    r.font.size = Pt(30)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("20 LAKARAN INTERFACE — GAYA STORYBOARD")
    r2.font.name = "Arial"
    r2.font.bold = True
    r2.font.size = Pt(22)
    r2.font.color.rgb = RGBColor(75, 75, 75)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_after = Pt(15)
    r3 = p3.add_run("Setiap storyboard menunjukkan:  INPUT  →  PROSES  →  POPUP  →  OUTPUT")
    r3.font.name = "Arial"
    r3.font.size = Pt(13)

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row, label in zip(table.rows, ["Nama Ahli 1", "Nama Ahli 2", "Nama Ahli 3", "Kursus / Kelas"]):
        row.cells[0].text = label
        row.cells[1].text = "____________________________________________"
        row.cells[0].paragraphs[0].runs[0].bold = True
        set_cell_shading(row.cells[0], "E7E5E0")
        for cell in row.cells:
            set_cell_margins(cell, 120, 150, 120, 150)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(16)
    p4.add_run("Lakaran awal / low-fidelity wireframe — bukan reka bentuk akhir").italic = True
    doc.add_page_break()

    # Allocation and index
    centered_heading(doc, "PEMBAHAGIAN TUGAS KUMPULAN", "Jumlah tepat: 20 interface")
    allocation = doc.add_table(rows=1, cols=4)
    allocation.style = "Table Grid"
    allocation.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(["Ahli", "Interface", "Bilangan", "Tugas tambahan"]):
        allocation.rows[0].cells[i].text = value
        set_cell_shading(allocation.rows[0].cells[i], "D6D3CD")
        allocation.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    rows = [
        ("Ahli 1", "01–07: Akaun, dashboard & kursus", "7", "Tetapkan gaya lakaran dan label."),
        ("Ahli 2", "08–14: Kursus, pembelajaran & tugasan", "7", "Semak aliran input hingga output."),
        ("Ahli 3", "15–20: Kuiz & komunikasi", "6", "Gabung dokumen dan pembentangan."),
    ]
    for values in rows:
        cells = allocation.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = value
            set_cell_margins(cells[i], 90, 100, 90, 100)

    doc.add_paragraph()
    index = doc.add_table(rows=1, cols=4)
    index.style = "Table Grid"
    index.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(["Bil.", "Interface", "Modul", "Urutan storyboard"]):
        index.rows[0].cells[i].text = value
        set_cell_shading(index.rows[0].cells[i], "E7E5E0")
        index.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for number, item in enumerate(SCREENS, 1):
        cells = index.add_row().cells
        values = (f"{number:02d}", item[0], item[1], "Input → Proses → Popup → Output")
        for i, value in enumerate(values):
            cells[i].text = value
            set_cell_margins(cells[i], 35, 70, 35, 70)
    doc.add_page_break()

    # One storyboard per interface.
    for number, (item, path) in enumerate(zip(SCREENS, paths), 1):
        title, module, fields, action, popup, output, _ = item
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"INTERFACE {number:02d}  |  {title.upper()}")
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(17)
        r2 = p.add_run(f"    Modul: {module}")
        r2.font.name = "Arial"
        r2.font.size = Pt(9)
        r2.font.color.rgb = RGBColor(90, 90, 90)
        picture_p = doc.add_paragraph()
        picture_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_p.paragraph_format.space_after = Pt(3)
        picture_p.add_run().add_picture(str(path), width=Inches(10.72), height=Inches(6.0))

        notes = doc.add_table(rows=1, cols=4)
        notes.style = "Table Grid"
        notes.alignment = WD_TABLE_ALIGNMENT.CENTER
        values = [
            "INPUT\n" + ", ".join(fields),
            "PROSES\n" + action,
            "POPUP\n“" + popup + "”",
            "OUTPUT\n" + output,
        ]
        for i, value in enumerate(values):
            notes.cell(0, i).text = value
            notes.cell(0, i).paragraphs[0].runs[0].font.size = Pt(8)
            notes.cell(0, i).paragraphs[0].runs[0].font.name = "Arial"
            set_cell_margins(notes.cell(0, i), 45, 70, 45, 70)
            set_cell_shading(notes.cell(0, i), "F4F2EC")
        if number < len(SCREENS):
            doc.add_page_break()

    for section in doc.sections:
        setup_section(section)
        add_page_number(section.footer.paragraphs[0])

    props = doc.core_properties
    props.title = "Student Learning Hub — 20 Storyboard Interface Sketches"
    props.subject = "Input, process, popup and output storyboard sketches"
    props.author = "Student Learning Hub Group"
    props.keywords = "storyboard, sketch, wireframe, student learning hub, interface"
    doc.save(OUTPUT)


paths = [draw_storyboard(i, item) for i, item in enumerate(SCREENS, 1)]
build_document(paths)
print(f"CREATED={OUTPUT}")
print(f"INTERFACES={len(paths)}")
print(f"PREVIEWS={PREVIEW_DIR}")
print(f"SIZE_BYTES={OUTPUT.stat().st_size}")
