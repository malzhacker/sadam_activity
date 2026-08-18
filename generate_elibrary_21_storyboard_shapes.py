"""Create 21 e-Library storyboard pages from basic editable Word shapes.

Each interface page contains four storyboard frames only:
Input, Process, Popup and Output. No explanatory notes section is added.
"""

from pathlib import Path
from textwrap import wrap

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

import generate_22_interfaces_doc as shapes
from generate_elibrary_21_word_shapes import SCREENS, TEAM

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "e_Library_System_21_Storyboard_Editable_Shapes.docx"
FONT = "Calibri"

box = shapes.box
label = shapes.label
circle = shapes.circle
rule = shapes.rule

# Four large storyboard frames on an A4 portrait page.
PANELS = [(22, 54), (310, 54), (22, 414), (310, 414)]
PANEL_W = 263
PANEL_H = 330
SCREEN_PAD = 10
SCREEN_Y_PAD = 34
SCREEN_W = 243
SCREEN_H = 270


def paragraph(doc, text, size=11, bold=False, align="left", before=0, after=4,
              color=(34, 34, 34)):
    p = doc.add_paragraph()
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT,
                   "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    return p


def set_cell(cell, text, bold=False, fill=None, size=8.5):
    cell.text = text
    for run in cell.paragraphs[0].runs:
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
    if fill:
        shapes.shading(cell, fill)
    shapes.cell_margins(cell, 55, 75, 55, 75)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def cover_page(doc):
    paragraph(doc, "DEPARTMENT OF INFORMATION AND COMMUNICATION TECHNOLOGY",
              11, True, "center", 46, 8)
    paragraph(doc, "DFC40343: SYSTEM ANALYSIS AND DESIGN FUNDAMENTALS",
              11, True, "center", 0, 30)
    paragraph(doc, "PROBLEM BASED ASSIGNMENT 1", 18, True, "center", 0, 32)
    paragraph(doc, "e-LIBRARY SYSTEM", 22, True, "center", 0, 5)
    paragraph(doc, "21 STORYBOARD INTERFACE SKETCHES", 13, False, "center", 0, 24,
              (80, 80, 80))

    project = doc.add_table(rows=5, cols=2)
    project.style = "Table Grid"
    project.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row, values in zip(project.rows, [
            ("PROJECT", "e-Library System"),
            ("LECTURER", "Pn Murniyati Binti Abdul"),
            ("CLASS", "DIT4-S1"),
            ("SUBMISSION", "19 July 2026"),
            ("DOCUMENT", "21 Storyboard Interfaces")]):
        set_cell(row.cells[0], values[0], True, "E5E5E5", 9)
        set_cell(row.cells[1], values[1], False, None, 9)

    paragraph(doc, "GROUP DIRECTORY", 11, True, "left", 14, 4)
    team_table = doc.add_table(rows=1, cols=3)
    team_table.style = "Table Grid"
    team_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(["NO.", "STUDENT NAME", "MATRIX NUMBER"]):
        set_cell(team_table.rows[0].cells[i], value, True, "D8D8D8", 8.5)
    for number, (name, matrix, _) in enumerate(TEAM, 1):
        cells = team_table.add_row().cells
        for i, value in enumerate((str(number), name, matrix)):
            set_cell(cells[i], value, False, "F7F7F7" if number % 2 == 0 else None, 8.2)

    paragraph(doc,
              "All interface sketches are made from basic Microsoft Word shapes: rectangle, "
              "rounded rectangle, oval, line, arrow and text box. No picture is embedded.",
              9, False, "center", 18, 0, (90, 90, 90))
    doc.add_page_break()


def index_page(doc):
    paragraph(doc, "e-LIBRARY STORYBOARD INDEX", 17, True, "left", 0, 3)
    paragraph(doc, "Exactly 21 interfaces based on the supplied CD, DFD and FDD.",
              9.5, False, "left", 0, 10, (90, 90, 90))
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(["ID", "INTERFACE", "USER", "FUNCTION"]):
        set_cell(table.rows[0].cells[i], value, True, "D8D8D8", 8.5)
    for number, screen in enumerate(SCREENS, 1):
        cells = table.add_row().cells
        for i, value in enumerate((f"EL-{number:02d}", screen[0], screen[1], screen[2])):
            set_cell(cells[i], value, False, "F7F7F7" if number % 2 == 0 else None, 8.1)
    doc.add_page_break()


def storyboard_frame(p, x, y, number, title):
    box(p, x, y, PANEL_W, PANEL_H, "", "FFFFFF", "333333", line_width=1.4)
    box(p, x, y, PANEL_W, 25, "", "E7E7E7", "333333")
    circle(p, x + 7, y + 4, 17, str(number), "FFFFFF", "555555", 7.5, True)
    label(p, x + 30, y + 5, PANEL_W - 38, 16, title, 8.5, True, "left")


def app_window(p, x, y, title):
    sx, sy = x + SCREEN_PAD, y + SCREEN_Y_PAD
    box(p, sx, sy, SCREEN_W, SCREEN_H, "", "FFFFFF", "666666")
    box(p, sx, sy, SCREEN_W, 22, "", "EEEEEE", "666666")
    label(p, sx + 7, sy + 4, 155, 14, "e-Library", 7.2, True, "left")
    label(p, sx + 166, sy + 4, 69, 14, title, 5.8, False, "right", "666666")
    box(p, sx, sy + 22, 47, SCREEN_H - 22, "", "F4F4F4", "999999")
    label(p, sx + 5, sy + 35, 37, 130,
          "HOME\n\nMEMBER\n\nBOOK\n\nLOAN\n\nFINE\n\nREPORT",
          5.8, False, "center", "555555")
    return sx + 56, sy + 31, SCREEN_W - 65, SCREEN_H - 40


def field(p, x, y, width, text, select=False):
    label(p, x, y, width, 12, text, 6.4, False, "left", "444444")
    box(p, x, y + 13, width, 20, "Select  v" if select else "Enter value",
        "FFFFFF", "888888", 6.2, False, "left", color="888888")


def input_scene(p, x, y, screen):
    cx, cy, cw, _ = app_window(p, x, y, "INPUT")
    label(p, cx, cy, cw, 15, screen[0], 8, True, "left")
    for index, item in enumerate(screen[3][:4]):
        fy = cy + 22 + index * 38
        select = any(word in item.lower() for word in
                     ("role", "status", "branch", "category", "format", "method",
                      "channel", "severity", "module", "condition", "period", "sort"))
        field(p, cx, fy, cw, item, select)
    button_y = min(cy + 181, cy + 22 + len(screen[3][:4]) * 38 + 5)
    box(p, cx, button_y, min(cw, 112), 24, screen[4], "E2E2E2", "444444",
        6.8, True, rounded=True)
    box(p, cx + min(cw, 112) + 7, button_y, 54, 24, "CLEAR", "FFFFFF", "777777",
        6.5, False, rounded=True)


def process_scene(p, x, y, screen):
    cx, cy, cw, _ = app_window(p, x, y, "PROCESS")
    label(p, cx, cy, cw, 15, "SYSTEM PROCESS", 8, True, "center")
    steps = ["Receive input", "Validate record", "Update database", "Prepare result"]
    center = cx + cw / 2
    for index, step in enumerate(steps):
        by = cy + 23 + index * 45
        box(p, center - 58, by, 116, 25, step, "FFFFFF", "555555",
            7, False, rounded=True)
        if index < len(steps) - 1:
            rule(p, center, by + 26, 1, 18, arrow=True, color="555555", width=1.0)
    # The boxes themselves show the process; no separate notes block is added.


def popup_scene(p, x, y, screen):
    cx, cy, cw, _ = app_window(p, x, y, "POPUP")
    # Simple background rows make the dialog visibly float over a working screen.
    for index in range(5):
        box(p, cx, cy + 12 + index * 31, cw, 20, "", "F8F8F8", "BBBBBB", dashed=True)
    modal_x, modal_y = cx + 12, cy + 55
    box(p, modal_x, modal_y, cw - 24, 120, "", "FFFFFF", "222222", line_width=1.5)
    box(p, modal_x, modal_y, cw - 24, 21, "SYSTEM MESSAGE", "E5E5E5", "222222",
        6.5, True, "left")
    circle(p, modal_x + 10, modal_y + 34, 25, "!", "FFFFFF", "555555", 9, True)
    message = "\n".join(wrap(screen[6], width=27))
    label(p, modal_x + 42, modal_y + 28, cw - 76, 54, message, 6.4, False, "left")
    if screen[8] == "confirm":
        box(p, modal_x + 23, modal_y + 87, 55, 22, "CANCEL", "FFFFFF", "666666",
            6.2, False, rounded=True)
        box(p, modal_x + 86, modal_y + 87, 55, 22, "CONFIRM", "E2E2E2", "444444",
            6.2, True, rounded=True)
    else:
        box(p, modal_x + cw - 84, modal_y + 87, 48, 22, "OK", "E2E2E2", "444444",
            6.5, True, rounded=True)


def output_table(p, cx, cy, cw, screen):
    widths = [int(cw * .48), int(cw * .25), cw - int(cw * .48) - int(cw * .25)]
    labels = ["Record", "Date", "Status"]
    x = cx
    for text, width in zip(labels, widths):
        box(p, x, cy, width, 21, text, "E5E5E5", "555555", 6.3, True, "left")
        x += width
    for row in range(4):
        x = cx
        values = [f"Item {row + 1}", "19 Jul", "Active" if row < 3 else "Closed"]
        for text, width in zip(values, widths):
            box(p, x, cy + 21 + row * 24, width, 24, text, "FFFFFF", "AAAAAA",
                6.1, False, "left")
            x += width


def output_scene(p, x, y, screen):
    cx, cy, cw, _ = app_window(p, x, y, "OUTPUT")
    label(p, cx, cy, cw, 15, "SYSTEM OUTPUT", 8, True, "left")
    kind = screen[8]
    if kind == "dashboard":
        cards = [("Members", "1,248"), ("Loans", "386"), ("Overdue", "27"), ("Fines", "RM820")]
        for index, (name, value) in enumerate(cards):
            col, row = index % 2, index // 2
            bx = cx + col * (cw / 2 + 2)
            by = cy + 22 + row * 57
            box(p, bx, by, cw / 2 - 4, 49, "", "FFFFFF", "777777", rounded=True)
            label(p, bx + 6, by + 5, cw / 2 - 16, 12, name, 6, False, "left", "666666")
            label(p, bx + 6, by + 19, cw / 2 - 16, 23, value, 10, True, "left")
        label(p, cx, cy + 145, cw, 14, "Borrowing trend", 6.8, True, "left")
        for index, height in enumerate([18, 32, 24, 45, 38]):
            box(p, cx + 14 + index * 29, cy + 204 - height, 16, height, "", "DDDDDD", "777777")
    elif kind in ("list", "report", "logs", "catalog"):
        output_table(p, cx, cy + 22, cw, screen)
        box(p, cx, cy + 136, 72, 22, "OPEN", "E2E2E2", "555555", 6.5, True, rounded=True)
        box(p, cx + 80, cy + 136, 72, 22, "EXPORT", "FFFFFF", "777777", 6.5, False, rounded=True)
    elif kind == "calculator":
        circle(p, cx + 5, cy + 27, 75, "RM\n3.00", "FFFFFF", "444444", 11, True)
        rows = [("Days overdue", "3"), ("Daily rate", "RM 1.00"),
                ("Adjustment", "RM 0.00"), ("Final fine", "RM 3.00")]
        for index, (name, value) in enumerate(rows):
            ry = cy + 29 + index * 34
            label(p, cx + 91, ry, 78, 13, name, 6.2, False, "left", "666666")
            box(p, cx + 91, ry + 14, cw - 92, 18, value, "FFFFFF", "999999",
                6.5, True, "left")
    else:
        box(p, cx, cy + 22, cw, 140, "", "FFFFFF", "777777")
        label(p, cx + 8, cy + 29, cw - 16, 16, "e-LIBRARY RECEIPT / RECORD", 7.2, True, "center")
        for index, (name, value) in enumerate([
                ("Reference", "EL-2026-0182"), ("Status", "SUCCESSFUL"),
                ("Date", "19 July 2026"), ("User", screen[1].split(" /")[0])]):
            ry = cy + 54 + index * 24
            label(p, cx + 10, ry, 64, 14, name, 6.2, False, "left", "666666")
            label(p, cx + 79, ry, cw - 89, 14, value, 6.6, True, "left")
            rule(p, cx + 10, ry + 16, cw - 20, color="BBBBBB")
    # The receipt, cards or table are the visual output; no notes block follows it.


def connect_panels(p):
    rule(p, 287, 215, 20, 1, arrow=True, color="555555", width=1.1)
    rule(p, 441, 389, 1, 20, arrow=True, color="555555", width=1.1)
    rule(p, 287, 575, 20, 1, arrow=True, color="555555", width=1.1)


def draw_page(p, number, screen):
    shapes.CURRENT_PAGE[0] = number
    label(p, 22, 16, 430, 24, f"EL-{number:02d}  {screen[0]}", 13, True, "left")
    label(p, 457, 19, 116, 18, screen[2], 7.5, False, "right", "555555")

    titles = ["INPUT", "PROCESS", "POPUP MESSAGE", "OUTPUT"]
    for index, ((x, y), title) in enumerate(zip(PANELS, titles), 1):
        storyboard_frame(p, x, y, index, title)
    input_scene(p, *PANELS[0], screen)
    process_scene(p, *PANELS[1], screen)
    popup_scene(p, *PANELS[2], screen)
    output_scene(p, *PANELS[3], screen)
    connect_panels(p)


def build():
    shapes.SHAPE_RECTS.clear()
    doc = Document()
    shapes.setup_section(doc.sections[0])
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"].font.size = Pt(10)
    cover_page(doc)
    index_page(doc)
    for number, screen in enumerate(SCREENS, 1):
        p = shapes.anchor_paragraph(doc)
        draw_page(p, number, screen)
        if number < len(SCREENS):
            doc.add_page_break()
    for section in doc.sections:
        shapes.setup_section(section)
    props = doc.core_properties
    props.title = "e-Library System - 21 Storyboard Interfaces"
    props.subject = "Four-frame storyboard sketches made from editable basic Word shapes"
    props.author = "Muhammad Haziq, Muhammad Danish and Seri Maisarah"
    props.keywords = "e-Library, storyboard, editable Word shapes, 21 interfaces"
    doc.save(OUTPUT)
    print(f"CREATED={OUTPUT.name}")
    print(f"INTERFACES={len(SCREENS)}")
    print(f"SHAPES={len(shapes.SHAPE_RECTS)}")
    print(f"SIZE_BYTES={OUTPUT.stat().st_size}")


if __name__ == "__main__":
    build()
